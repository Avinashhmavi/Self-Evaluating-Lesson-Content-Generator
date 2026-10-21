import io, json, re
from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
import config, graph, memory
from rubric import RUBRIC

app = FastAPI(title="Lesson Gate")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"], allow_methods=["*"], allow_headers=["*"])


class RunReq(BaseModel):
    topic: str = Field(min_length=2, max_length=120)
    inject_error: bool = False


def sse(event, data):
    return f"event: {event}\ndata: {json.dumps(data)}\n\n"


@app.post("/api/run")
def run(req: RunReq):
    def gen():
        try:
            for node, upd in graph.stream(req.topic.strip(), req.inject_error):
                if node == "plan_topic":
                    upd = {"profile": {k: v for k, v in upd["profile"].items() if k != "sources"}}
                yield sse("node", {"node": node, "update": upd})
            yield sse("done", {})
        except Exception as e:
            yield sse("error", {"message": str(e)})
    return StreamingResponse(gen(), media_type="text/event-stream", headers={"Cache-Control": "no-cache"})


@app.get("/api/rubric")
def rubric():
    return {"checks": [dict(id=r[0], name=r[1], kind=r[2], tech=r[3], condition=r[4], hint=r[5]) for r in RUBRIC],
            "max_attempts": config.MAX_ATTEMPTS, "generator": config.GEN_MODEL, "judges": config.JUDGE_MODELS}


@app.get("/api/memory")
def mem():
    return memory.stats()


@app.get("/api/runs/{rid}")
def get_run(rid: int):
    r = memory.get_run(rid)
    if not r:
        raise HTTPException(404)
    return r


@app.get("/api/runs/{rid}/docx")
def get_docx(rid: int):
    r = memory.get_run(rid)
    if not r:
        raise HTTPException(404)
    doc = Document()
    doc.add_heading(r["topic"], 0)
    for block in r["lesson"].split("\n\n"):
        block = block.strip()
        if block.startswith("## "):
            doc.add_heading(block[3:], 1)
        elif block:
            doc.add_paragraph(re.sub(r"[*_`]", "", block).replace("\n", " "))
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    name = re.sub(r"[^a-z0-9]+", "-", r["topic"].lower()).strip("-")[:60] or "lesson"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f'attachment; filename="{name}.docx"'})


@app.post("/api/proposals/{pid}/{action}")
def proposal(pid: int, action: str):
    if action not in ("approve", "reject"):
        raise HTTPException(400)
    memory.set_proposal(pid, "approved" if action == "approve" else "rejected")
    return {"ok": True}
