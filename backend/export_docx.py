import sys
from docx import Document
import memory

rid = int(sys.argv[1])
run = memory.get_run(rid)
body = "# Introduction to RAG\n\n" + run["lesson"]
open(f"docs/lesson_rag.md", "w").write(body + "\n")
doc = Document()
for block in body.split("\n\n"):
    block = block.strip()
    if block.startswith("# "):
        doc.add_heading(block[2:], 0)
    elif block.startswith("## "):
        doc.add_heading(block[3:], 1)
    elif block:
        doc.add_paragraph(block.replace("\n", " "))
doc.save("docs/lesson_rag.docx")
print("run", rid, run["status"], "attempts", run["attempts"])
